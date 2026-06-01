"""
excel_to_json.py
────────────────
Reads semiconductor_data.xlsx and writes data.json for the map.

Usage:
    python excel_to_json.py                          # uses default paths
    python excel_to_json.py path/to/data.xlsx        # custom Excel path

Sheets read:
    FAB_COUNT_BY_COUNTRY  → data["fab_count_by_country"]
    FABS                  → data["fabs"]
    MILESTONES            → data["milestones"]

Companies and coordinates are kept from the existing data.json
(they don't have a sheet since they rarely change).
"""

import json
import sys
import os
from openpyxl import load_workbook

# ── paths ─────────────────────────────────────────────────────────────────────
SCRIPT_DIR  = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH  = sys.argv[1] if len(sys.argv) > 1 else os.path.join(SCRIPT_DIR, "semiconductor_data.xlsx")
DATA_JSON   = os.path.join(SCRIPT_DIR, "data.json")
YEARS       = [2020, 2021, 2022, 2023, 2024, 2025, 2026]


def load_existing_data():
    with open(DATA_JSON) as f:
        return json.load(f)


def parse_fab_count(wb):
    """
    Sheet: FAB_COUNT_BY_COUNTRY
    Layout:
      Row 3: merged year headers
      Row 4: Active | Under Constr. sub-headers  (cols 3,4 = 2020; 5,6 = 2021 …)
      Row 5+: Country | Source | active | constr | active | constr | …
      Last row: TOTAL (skip)
    """
    ws = wb["FAB_COUNT_BY_COUNTRY"]
    result = {str(y): {} for y in YEARS}

    for row in ws.iter_rows(min_row=5, values_only=True):
        country = row[0]
        if not country or str(country).strip().upper() == "TOTAL":
            continue
        country = str(country).strip()
        col_offset = 2  # 0-indexed: col 0=country, 1=source, then pairs
        for yi, year in enumerate(YEARS):
            ci = col_offset + yi * 2
            active = row[ci] if ci < len(row) else None
            constr = row[ci + 1] if ci + 1 < len(row) else None
            try:
                active = int(active) if active and str(active) != "-" else 0
            except (ValueError, TypeError):
                active = 0
            try:
                constr = int(constr) if constr and str(constr) != "-" else 0
            except (ValueError, TypeError):
                constr = 0
            result[str(year)][country] = {"active": active, "construction": constr}

    return result


def parse_fabs(wb):
    """
    Sheet: FABS
    Row 3: headers
    Row 4+: fab_id | company | cat | loc | name | country | 2020..2026
    """
    ws = wb["FABS"]
    fabs = []

    for row in ws.iter_rows(min_row=4, values_only=True):
        fab_id = row[0]
        if not fab_id or str(fab_id).startswith("#"):
            continue
        fab_id  = str(fab_id).strip()
        company = str(row[1]).strip() if row[1] else ""
        loc     = str(row[3]).strip() if row[3] else ""
        name    = str(row[4]).strip() if row[4] else ""

        years = {}
        for yi, year in enumerate(YEARS):
            ci = 6 + yi  # 0-indexed: cols 0-5 = metadata, 6-12 = years
            val = row[ci] if ci < len(row) else None
            if val and str(val).strip():
                years[str(year)] = str(val).strip()

        fabs.append({
            "id":      fab_id,
            "company": company,
            "loc":     loc,
            "name":    name,
            "years":   years,
        })

    return fabs


def parse_milestones(wb):
    """
    Sheet: MILESTONES
    Row 2: headers
    Row 3+: year | event | notes | source
    """
    ws = wb["MILESTONES"]
    milestones = {str(y): [] for y in YEARS}

    for row in ws.iter_rows(min_row=3, values_only=True):
        year = row[0]
        event = row[1]
        if not year or not event:
            continue
        try:
            year_key = str(int(float(str(year))))
        except (ValueError, TypeError):
            continue
        if year_key in milestones:
            milestones[year_key].append(str(event).strip())

    return milestones


def main():
    print(f"📂 Reading Excel: {EXCEL_PATH}")
    if not os.path.exists(EXCEL_PATH):
        print(f"❌ File not found: {EXCEL_PATH}")
        sys.exit(1)

    wb = load_workbook(EXCEL_PATH, data_only=True)
    print(f"   Sheets found: {wb.sheetnames}")

    # Load existing data.json to keep companies + coords
    data = load_existing_data()
    print(f"   Existing companies: {len(data['companies'])}")
    print(f"   Existing coords:    {len(data['coords'])}")

    # Parse sheets
    fab_count = parse_fab_count(wb)
    fabs      = parse_fabs(wb)
    milestones = parse_milestones(wb)

    print(f"   Parsed fab_count: {len(fab_count)} years")
    print(f"   Parsed fabs:      {len(fabs)}")
    total_ms = sum(len(v) for v in milestones.values())
    print(f"   Parsed milestones:{total_ms} events")

    # Update data
    data["fab_count_by_country"] = fab_count
    data["fabs"]                 = fabs
    data["milestones"]           = milestones

    # Save
    with open(DATA_JSON, "w") as f:
        json.dump(data, f, indent=2)

    print(f"\n✅ data.json updated → {DATA_JSON}")
    print("   Next step: git add . && git commit -m 'update data' && git push")


if __name__ == "__main__":
    main()
